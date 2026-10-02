"""Spreadsheet safety must never contaminate raw training/checkpoint data."""

import csv
from datetime import UTC, datetime
from io import BytesIO, StringIO

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from sentiment_prep.errors import ValidationError
from sentiment_prep.export.csv_export import to_checkpoint_csv, to_csv
from sentiment_prep.export.excel_export import to_excel
from sentiment_prep.export.parquet_export import EXPORT_SCHEMA, csv_to_parquet, to_parquet
from sentiment_prep.export.rows import ROW_COLUMNS, bundle_rows
from sentiment_prep.models import Dataset, DatasetBundle, Record
from sentiment_prep.storage.checkpoints import LocalCheckpointStore, checkpoint_bundle


def make_bundle(records):
    dataset = Dataset(records=records, source_type="csv")
    return DatasetBundle(dataset_id="export-test", original=dataset, processed=dataset)


def parquet_rows(data):
    return pq.read_table(BytesIO(data))


@pytest.mark.parametrize("prefix", ["=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)"])
@pytest.mark.parametrize("leading", ["", " ", "\t", "\r", "\n", "\u00a0", "\u200b", "\ufeff"])
def test_spreadsheet_exports_neutralize_every_untrusted_string_column(prefix, leading):
    text = leading + prefix
    bundle = make_bundle(
        [
            Record(
                id=text,
                text=text,
                label=text,
                comprehend_label=text,
                tokens=[text],
                source_type="csv",
            )
        ]
    )
    raw = bundle_rows(bundle)[0]
    csv_row = next(csv.DictReader(StringIO(to_csv(bundle).decode("utf-8-sig"))))
    for name in ("id", "label", "comprehend_label", "original_text", "processed_text"):
        assert csv_row[name] == "'" + text
    workbook = load_workbook(BytesIO(to_excel(bundle)), data_only=False)
    for name, cell in zip(ROW_COLUMNS, workbook["data"][2], strict=True):
        if isinstance(raw[name], str):
            # XML readers normalize literal CR line endings; spreadsheet text remains inert.
            assert cell.value == raw[name].replace("\r\n", "\n").replace("\r", "\n")
            assert cell.data_type == "s", name
    assert parquet_rows(to_parquet(bundle)).to_pylist() == [raw]
    assert parquet_rows(csv_to_parquet(to_checkpoint_csv(bundle))).to_pylist() == [raw]
    assert bundle.original.records[0].text == text


@pytest.mark.parametrize("control", ["\x00", "\x01", "\x1f", "\x7f", "\x85", "\u2060"])
def test_csv_control_prefixes_are_neutralized_without_changing_raw_checkpoints(control):
    text = control + "=1+1"
    bundle = make_bundle([Record(id="record", text=text, source_type="csv")])
    assert (
        next(csv.DictReader(StringIO(to_csv(bundle).decode("utf-8-sig"))))["original_text"]
        == "'" + text
    )
    assert (
        parquet_rows(csv_to_parquet(to_checkpoint_csv(bundle))).to_pylist()[0]["original_text"]
        == text
    )


def test_checkpoint_store_uses_raw_csv(tmp_path):
    bundle = make_bundle([Record(id="00123", text="=1+1", source_type="csv")])
    store = LocalCheckpointStore(tmp_path)
    checkpoint_bundle(store, bundle, "collected")
    raw = store.read(bundle.dataset_id, "collected", "csv")
    assert raw == to_checkpoint_csv(bundle)
    assert parquet_rows(csv_to_parquet(raw)).to_pylist()[0]["original_text"] == "=1+1"


def test_explicit_schema_preserves_ids_text_nulls_and_empty_strings():
    records = [
        Record(
            id="00123",
            text="12345",
            source_type="csv",
            label="0007",
            label_confidence=0.0,
            created_at=datetime(2026, 1, 1, tzinfo=UTC),
        ),
        Record(id="999999999999999999999999999999999999", text="", source_type="csv", label=""),
        Record(id="3", text="NA", source_type="csv"),
    ]
    bundle = make_bundle(records)
    direct = parquet_rows(to_parquet(bundle))
    converted = parquet_rows(csv_to_parquet(to_checkpoint_csv(bundle)))
    assert direct.schema == converted.schema == EXPORT_SCHEMA
    assert direct.to_pylist() == converted.to_pylist() == bundle_rows(bundle)
    assert direct.schema.field("id").type == pa.string()
    assert direct.schema.field("original_text").type == pa.string()
    assert direct.schema.field("created_at").type == pa.string()  # ISO-8601, with timezone retained
    assert direct.schema.field("label_confidence").type == pa.float64()
    assert direct.schema.field("comprehend_confidence").type == pa.float64()
    assert direct.to_pylist()[1]["original_text"] == ""
    assert direct.to_pylist()[2]["label"] is None


def test_empty_export_retains_schema_in_both_paths():
    bundle = make_bundle([])
    direct = parquet_rows(to_parquet(bundle))
    converted = parquet_rows(csv_to_parquet(to_checkpoint_csv(bundle)))
    assert direct.schema == converted.schema == EXPORT_SCHEMA
    assert direct.num_rows == converted.num_rows == 0


def test_legacy_numeric_looking_csv_checkpoint_still_uses_string_columns():
    bundle = make_bundle([Record(id="00123", text="12345", label="007", source_type="csv")])
    table = parquet_rows(csv_to_parquet(to_csv(bundle)))
    assert table.to_pylist()[0]["id"] == "00123"
    assert table.to_pylist()[0]["original_text"] == "12345"
    assert table.to_pylist()[0]["label"] == "007"


def test_excel_preserves_text_at_the_cell_limit():
    text = "x" * 32767
    workbook = load_workbook(
        BytesIO(to_excel(make_bundle([Record(id="1", text=text, source_type="csv")])))
    )
    assert workbook["data"].cell(2, ROW_COLUMNS.index("original_text") + 1).value == text


@pytest.mark.parametrize("field", ["text", "id", "label"])
def test_excel_refuses_silent_truncation_but_lossless_formats_keep_the_value(field):
    text = "x" * 32768
    record = Record.model_validate({"id": "1", "text": "hello", "source_type": "csv", field: text})
    bundle = make_bundle([record])
    with pytest.raises(ValidationError, match=r"32,767.*row 2.*Download CSV or Parquet"):
        to_excel(bundle)
    column = "original_text" if field == "text" else field
    assert next(csv.DictReader(StringIO(to_csv(bundle).decode("utf-8-sig"))))[column] == text
    assert parquet_rows(to_parquet(bundle)).to_pylist()[0][column] == text


@pytest.mark.parametrize("field", ["id", "label", "text"])
@pytest.mark.parametrize(
    "character", ["\x00", "\x01", "\x08", "\x0b", "\x0c", "\x0e", "\x1f", "\ufffe", "\uffff"]
)
def test_excel_checks_all_string_columns_for_invalid_xml_without_mutating_records(field, character):
    text = f"hello{character}world"
    bundle = make_bundle(
        [Record.model_validate({"id": "1", "text": "hello", "source_type": "csv", field: text})]
    )
    with pytest.raises(ValidationError, match=f"U\\+{ord(character):04X}"):
        to_excel(bundle)
    assert getattr(bundle.original.records[0], field) == text
    column = "original_text" if field == "text" else field
    assert parquet_rows(to_parquet(bundle)).to_pylist()[0][column] == text


@pytest.mark.parametrize(
    "text", ["hello\tworld", "hello\nworld", "café 中文 🙂", "hello\x7fworld", "hello\x85world"]
)
def test_excel_keeps_valid_xml_whitespace_and_unicode(text):
    workbook = load_workbook(
        BytesIO(to_excel(make_bundle([Record(id="1", text=text, source_type="csv")])))
    )
    assert workbook["data"].cell(2, ROW_COLUMNS.index("original_text") + 1).value == text
