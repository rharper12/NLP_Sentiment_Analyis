import json

from openpyxl import load_workbook

from sentiment_prep.analysis.comprehend_scorer import ComprehendScorer
from sentiment_prep.analysis.embeddings import EmbeddingDrift, cosine
from sentiment_prep.analysis.metrics import compute_metrics
from sentiment_prep.export.csv_export import to_csv
from sentiment_prep.export.excel_export import to_excel
from sentiment_prep.models import DatasetBundle, ImpactReport
from sentiment_prep.preprocessing import DEFAULT_ORDER, STEP_REGISTRY, Pipeline
from sentiment_prep.report import render_report
from sentiment_prep.storage.s3_store import S3Store
from tests.conftest import FakeBedrock, FakeComprehend, make_dataset


def processed_bundle() -> DatasetBundle:
    ds = make_dataset()
    processed, steps = Pipeline([STEP_REGISTRY[n]() for n in DEFAULT_ORDER]).run(ds)
    return DatasetBundle(
        dataset_id="abc",
        original=ds,
        processed=processed,
        applied_steps=DEFAULT_ORDER,
        report=ImpactReport(steps=steps),
    )


def test_metrics_shape():
    m = compute_metrics(make_dataset())
    assert m.record_count == 8
    assert sum(m.length_histogram.values()) == 8
    assert 0 < m.type_token_ratio <= 1


def test_comprehend_never_pays_twice_for_the_same_text():
    """Every document is sent once however often it is scored, because each send is billed.

    Asserted as "no further calls", not as an exact call count: how many documents fit in a batch
    is Comprehend's business, and a test that pins it breaks when the limit changes without
    anything actually being wrong.
    """
    ds = make_dataset([f"unique text {i}" for i in range(60)])
    fake = FakeComprehend()
    scorer = ComprehendScorer(fake)

    first = scorer.label(ds)
    calls_after_first = fake.calls
    assert calls_after_first > 0
    assert len({r.text for r in ds.records}) == 60  # all distinct, so all had to be sent

    second = scorer.label(ds)
    assert fake.calls == calls_after_first  # nothing re-sent
    assert [x.label for x in second] == [x.label for x in first]


def test_comprehend_labels_carry_confidence():
    ds = make_dataset(["I loved it", "terrible", "meh"])
    labels = ComprehendScorer(FakeComprehend()).label(ds)
    assert [x.label for x in labels] == ["positive", "negative", "neutral"]
    assert labels[0].confidence == 0.85


def test_comprehend_compare_aligns_by_id():
    b = processed_bundle()
    cmp = ComprehendScorer(FakeComprehend()).compare(b.original, b.processed)
    assert 0 <= cmp.agreement <= 1
    assert sum(cmp.distribution_before.values()) == 8
    assert sum(cmp.distribution_after.values()) == 6


def test_embedding_drift_zero_for_identity():
    ds = make_dataset()
    assert EmbeddingDrift(FakeBedrock(), "m", 5).compute(ds, ds) == 0.0
    assert cosine([0.0, 0.0], [1.0, 1.0]) == 0.0


def test_exports_and_report(s3_bucket):
    b = processed_bundle()
    csv_bytes = to_csv(b)
    assert csv_bytes.startswith("\ufeff".encode()) and b"processed_text" in csv_bytes

    wb = load_workbook(filename=__import__("io").BytesIO(to_excel(b)))
    assert wb.sheetnames == ["data", "impact"]
    assert wb["impact"].max_row >= len(DEFAULT_ORDER) + 1

    md = render_report(b)
    assert "## Strengths and limitations" in md and "Remove stopwords" in md
    assert b"label_source" in csv_bytes

    client, bucket = s3_bucket
    uri = S3Store(bucket, "datasets", client).save(b)
    keys = [o["Key"] for o in client.list_objects_v2(Bucket=bucket)["Contents"]]
    assert any(k.endswith("dataset.parquet") for k in keys)
    manifest = json.loads(
        client.get_object(Bucket=bucket, Key=next(k for k in keys if k.endswith("manifest.json")))[
            "Body"
        ].read()
    )
    assert manifest["record_count_original"] == 8
    assert uri.startswith(f"s3://{bucket}/datasets/abc/")


def test_generated_explanations_never_receive_operator_timings_or_hints():
    from sentiment_prep.analysis.bedrock_explainer import BedrockExplainer

    bundle = processed_bundle()
    bundle.report.warnings = ["internal-only-configuration-hint"]
    fake = FakeBedrock()
    BedrockExplainer(fake, "model").explain(bundle.report, bundle.applied_steps)
    assert "duration_ms" not in fake.last_prompt
    assert "internal-only-configuration-hint" not in fake.last_prompt
    assert "sample_diffs" not in fake.last_prompt
