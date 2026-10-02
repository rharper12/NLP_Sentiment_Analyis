"""Large accepted uploads remain pageable and downloadable through the Lambda adapter."""

import asyncio
import base64
import csv
import io
import json
from types import SimpleNamespace
from urllib.parse import unquote, urlparse

import pytest
from fastapi.testclient import TestClient
from mangum import Mangum

from sentiment_prep.api import deps
from sentiment_prep.api.app import create_app
from sentiment_prep.api.downloads import DOWNLOAD_LINK_MEDIA_TYPE
from sentiment_prep.config import Settings, get_settings
from sentiment_prep.storage.checkpoints import LocalCheckpointStore
from sentiment_prep.storage.repository import InMemoryRepository
from sentiment_prep.storage.s3_store import S3Store


@pytest.fixture
def large_dataset(tmp_path):
    repo = InMemoryRepository()
    settings = Settings(
        _env_file=None, dedupe_enabled=False, comprehend_enabled=False, diagnostics=False
    )
    app = create_app()
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[deps.get_repository] = lambda: repo
    app.dependency_overrides[deps.get_checkpoint_store] = lambda: LocalCheckpointStore(tmp_path)
    client = TestClient(app)
    content = io.StringIO()
    writer = csv.writer(content)
    writer.writerow(["id", "text"])
    for i in range(500):
        writer.writerow([str(i), (f"WORD{i:04d} " * 875).strip()])
    raw = content.getvalue().encode()
    assert len(raw) < 4 * 1024 * 1024
    uploaded = client.post("/dataset/upload", files={"file": ("large.csv", raw)})
    assert uploaded.status_code == 200
    assert uploaded.json()["record_count"] == 500
    identifier = uploaded.json()["dataset_id"]
    response = client.post(f"/dataset/{identifier}/preprocess", json={"steps": ["lowercase"]})
    assert response.status_code == 200
    return app, client, identifier


def invoke(app, path, query=""):
    event = {
        "version": "2.0",
        "routeKey": "$default",
        "rawPath": path,
        "rawQueryString": query,
        "headers": {"host": "localhost", "x-api-key": "test-key"},
        "requestContext": {
            "http": {
                "method": "GET",
                "path": path,
                "sourceIp": "127.0.0.1",
                "protocol": "HTTP/1.1",
            },
            "stage": "$default",
        },
        "isBase64Encoded": False,
    }
    try:
        previous_loop = asyncio.get_event_loop()
    except RuntimeError:
        previous_loop = None
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    try:
        result = Mangum(app, lifespan="off")(event, SimpleNamespace(aws_request_id="response-test"))
    finally:
        loop.close()
        asyncio.set_event_loop(previous_loop)
    assert result["statusCode"] == 200
    assert len(json.dumps(result).encode()) < 6 * 1024 * 1024
    body = base64.b64decode(result["body"]) if result["isBase64Encoded"] else result["body"]
    return result, json.loads(body)


def test_records_pages_fit_lambda_and_cover_every_original_and_processed_row(large_dataset):
    app, _, identifier = large_dataset
    records = []
    while len(records) < 500:
        _, page = invoke(app, f"/dataset/{identifier}/records", f"offset={len(records)}&limit=1000")
        assert page["total"] == 500 and page["offset"] == len(records)
        assert 0 < len(page["items"]) < 500
        records.extend(page["items"])
    assert [r["original"]["id"] for r in records] == [str(i) for i in range(500)]
    assert all(r["processed"]["text"] == r["original"]["text"].lower() for r in records)


@pytest.mark.parametrize(
    "kind", ["export.csv", "export.xlsx", "export.parquet", "original.json", "report.md"]
)
def test_cloud_download_preserves_local_export_bytes_without_a_large_lambda_response(
    large_dataset, s3_bucket, monkeypatch, kind
):
    app, client, identifier = large_dataset
    path = f"/dataset/{identifier}/{kind}"
    local = client.get(path)
    assert local.status_code == 200
    if kind == "export.csv":
        assert len(local.content) > 6 * 1024 * 1024
    s3, bucket = s3_bucket
    monkeypatch.setattr(deps, "get_s3_store", lambda: S3Store(bucket, "datasets", s3))
    settings = Settings(_env_file=None, runtime="lambda", data_bucket=bucket, api_key="test-key")
    app.dependency_overrides[get_settings] = lambda: settings
    response, link = invoke(app, path)
    assert response["headers"]["content-type"] == DOWNLOAD_LINK_MEDIA_TYPE
    assert response["headers"]["cache-control"] == "no-store"
    assert link["expires_in"] == 300 and urlparse(link["url"]).scheme == "https"
    key = unquote(urlparse(link["url"]).path).lstrip("/")
    assert key.startswith("_downloads/")
    saved = s3.get_object(Bucket=bucket, Key=key)
    assert saved["ContentDisposition"] == local.headers["content-disposition"]
    content = saved["Body"].read()
    if kind != "export.xlsx":  # ZIP timestamps can change between invocations.
        assert content == local.content
    else:
        from openpyxl import load_workbook

        expected = load_workbook(io.BytesIO(local.content))
        actual = load_workbook(io.BytesIO(content))
        assert list(actual["data"].values) == list(expected["data"].values)


def test_excel_limit_is_an_actionable_api_error(tmp_path):
    app = create_app()
    app.dependency_overrides[deps.get_repository] = lambda: repo
    app.dependency_overrides[deps.get_checkpoint_store] = lambda: LocalCheckpointStore(tmp_path)
    repo = InMemoryRepository()
    client = TestClient(app)
    uploaded = client.post(
        "/dataset/upload", files={"file": ("long.csv", b"text\n" + b"x" * 40000)}
    )
    assert uploaded.status_code == 200
    identifier = uploaded.json()["dataset_id"]
    response = client.get(f"/dataset/{identifier}/export.xlsx")
    assert response.status_code == 400
    assert "Download CSV or Parquet" in response.json()["error"]
