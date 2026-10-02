"""Excel export: a ``data`` sheet plus an ``impact`` sheet with per-step statistics."""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sentiment_prep.errors import ValidationError
from sentiment_prep.export.rows import ROW_COLUMNS, bundle_rows
from sentiment_prep.models import DatasetBundle

IMPACT_COLUMNS = [
    "step_name",
    "records_in",
    "records_out",
    "vocab_before",
    "vocab_after",
    "avg_tokens_before",
    "avg_tokens_after",
    "duration_ms",
]
MAX_CELL_CHARACTERS = 32767
# XML 1.0 permits tabs, line breaks, and ordinary Unicode, but not these code points.
_INVALID_XML_CHARACTER = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def to_excel(bundle: DatasetBundle, *, diagnostics: bool = False) -> bytes:
    """Build the workbook in memory."""
    workbook = Workbook()
    data = workbook.active
    data.title = "data"
    _write_sheet(data, ROW_COLUMNS, [[row[c] for c in ROW_COLUMNS] for row in bundle_rows(bundle)])

    impact = workbook.create_sheet("impact")
    steps = bundle.report.steps if bundle.report else []
    columns = IMPACT_COLUMNS if diagnostics else [c for c in IMPACT_COLUMNS if c != "duration_ms"]
    _write_sheet(impact, columns, [[getattr(s, c) for c in columns] for s in steps])
    if bundle.report:
        impact.append([])
        impact.append(
            [
                "sentiment_agreement",
                bundle.report.sentiment.agreement if bundle.report.sentiment else None,
            ]
        )
        if bundle.report.sentiment:
            impact.append(
                ["sentiment_comparable_records", bundle.report.sentiment.comparable_records]
            )
            impact.append(["sentiment_shared_records", bundle.report.sentiment.shared_records])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_sheet(sheet: Worksheet, columns: list[str], rows: list[list[object]]) -> None:
    """Write a header row and body, then size and freeze the header."""
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row_number, row in enumerate(rows, start=2):
        for column, value in zip(columns, row, strict=True):
            if not isinstance(value, str):
                continue
            if len(value) > MAX_CELL_CHARACTERS:
                raise ValidationError(
                    f"Excel cannot store more than {MAX_CELL_CHARACTERS:,} characters per cell "
                    f"({sheet.title}, row {row_number}, {column}). "
                    "Download CSV or Parquet to preserve the complete text."
                )
            if invalid := _INVALID_XML_CHARACTER.search(value):
                raise ValidationError(
                    f"Excel cannot store character U+{ord(invalid[0]):04X} "
                    f"({sheet.title}, row {row_number}, {column}). "
                    "Download CSV or Parquet to preserve the complete text."
                )
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Dataset strings must never become executable formulas.
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = 40 if "text" in column else 16
    sheet.freeze_panes = "A2"
